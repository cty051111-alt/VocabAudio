"""Vocabulary parsing and audio composition. No UI dependency."""
from __future__ import annotations
import re, subprocess, tempfile, time, urllib.parse
from pathlib import Path

HEADERS = {'单词','英文单词','英语单词','词汇','word','words','vocabulary','英文','序号','音标','词性','中文释义','中文','释义','meaning','definition','默写','默写区','词形','词形拓展','词形拓展／搭配','拓展','搭配','复习'}
IPA_MARKERS = ('/', '[', 'ˈ', 'ˌ', 'ə', 'ɪ', 'ʊ', 'ŋ', 'æ', 'ɔ', 'ʒ', 'ð', 'θ')
WORD_RE = re.compile(r"^[A-Za-z][A-Za-z'-]{1,40}(?:\s+[A-Za-z][A-Za-z'-]{1,32}){0,3}$")
LINE_WORD = re.compile(r"^\s*(?:(?:\d{1,3})\s*[.、)\-]?\s+)?([A-Za-z][A-Za-z'-]{1,39}(?:\s+[A-Za-z][A-Za-z'-]{1,28}){0,2})(?=\s*(?:/|\[|\||\t| {2,}|（|\(|—|–|:|：|$))")


def normalise_word(value: str) -> str:
    value = value.strip().replace('’', "'").replace('‐','-').replace('‑','-')
    value = re.sub(r'^\s*(?:\d{1,3}[.、)\-]?\s*)', '', value)
    return re.sub(r'\s+', ' ', value).strip()


def valid_word(s: str) -> bool:
    w = normalise_word(s)
    return bool(WORD_RE.fullmatch(w)) and w.lower() not in HEADERS and not w.lower().startswith(('http:', 'https:'))


def deduplicate(words):
    seen = set(); out=[]
    for word in words:
        w=normalise_word(word)
        if valid_word(w) and w.casefold() not in seen:
            out.append(w); seen.add(w.casefold())
    return out


def row_word(cells):
    """Prefer vocabulary/headword column; support leading ordinal and British IPA."""
    arr=[re.sub(r'\s+', ' ', str(c or '')).strip() for c in cells]
    for cell in arr[:4]:
        w=normalise_word(cell)
        if valid_word(w) and len(w.split()) <= 3:
            return w
        if not cell or cell.lower() in HEADERS or re.fullmatch(r'\d+',cell):
            continue
        # One wide cell may include headword, IPA and Chinese translation.
        m=LINE_WORD.match(cell)
        if m and valid_word(m.group(1)):
            return normalise_word(m.group(1))
    return None


def from_lines(lines):
    words=[]
    for line in lines:
        s=line.strip()
        if not s or re.fullmatch(r'\d+', s) or s.lower() in HEADERS:
            continue
        # Normal columns: tab or two+ spaces, or pipe.
        fields=re.split(r'\t|\s*\|\s*| {2,}',s)
        found=row_word(fields)
        if found:
            words.append(found); continue
        m=LINE_WORD.match(s)
        if m: words.append(m.group(1))
    return deduplicate(words)


def read_docx(path):
    from docx import Document
    doc=Document(str(path)); collected=[]
    for table in doc.tables:
        for row in table.rows:
            v=row_word([cell.text for cell in row.cells])
            if v: collected.append(v)
    if collected:
        return deduplicate(collected), 'Word 表格'
    return from_lines(p.text for p in doc.paragraphs), 'Word 正文'


def read_pdf(path):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc=fitz.open(str(path)); fallback_lines=[]; table_words=[]
    # A vocabulary PDF often has a "word" column followed by an IPA column.
    # Inspect geometric column headers and only read cells from the headword column.
    for page in doc:
        all_lines=[]
        for block in page.get_text('dict').get('blocks', []):
            for line in block.get('lines', []):
                text=' '.join(span['text'] for span in line['spans']).strip()
                if text:
                    x0,y0,x1,y1=line['bbox'];all_lines.append((x0,y0,x1,y1,text))
                    fallback_lines.append(text)
        header = [t for t in all_lines if
                  ('word / phrase' in t[4].lower() or t[4].lower().strip() in ('word','words','english','英文单词','英语单词','单词'))]
        ipa = [t for t in all_lines if ('音标' in t[4] or 'ipa' == t[4].lower().strip())]
        if header and ipa:
            head=min(header,key=lambda x:x[1]); nextcol=min(ipa,key=lambda x:abs(x[1]-head[1]))
            if nextcol[0]>head[0]+12:
                lo=head[0]-4; hi=nextcol[0]-1
                entries=[]
                for x0,y0,x1,y1,text in all_lines:
                    if y0<=head[1]+5: continue
                    if lo<=x0<hi and valid_word(text) and text.lower() not in ('word / phrase','word','words'):
                        entries.append((y0,x0,normalise_word(text)))
                table_words.extend(w for _,__,w in sorted(entries))
    if table_words:
        return deduplicate(table_words),'PDF 英文词条列'
    return from_lines(fallback_lines),'PDF 文字层（请人工核对）'

def read_txt(path):
    raw=Path(path).read_text(encoding='utf-8-sig')
    return from_lines(raw.splitlines()), '文本行'


def read_file(path):
    path=Path(path)
    ext=path.suffix.lower()
    if ext=='.pdf': return read_pdf(path)
    if ext=='.docx': return read_docx(path)
    if ext in ('.txt','.csv'): return read_txt(path)
    if ext in ('.wps','.doc'):
        raise ValueError('旧版 .wps/.doc 文件不能可靠地直接提取。请在 WPS 中选择「另存为 .docx」或「导出 PDF」，再导入。')
    raise ValueError(f'不支持的文件类型：{ext}。请使用 PDF、DOCX 或 TXT。')


def youdao_url(word, accent='us'):
    return 'https://dict.youdao.com/dictvoice?' + urllib.parse.urlencode({'audio':word, 'type': '2' if accent=='us' else '1'})


def fetch_clip(word, dest, accent='us', timeout=18):
    import requests
    url=youdao_url(word,accent)
    response=requests.get(url, headers={'User-Agent':'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36','Referer':'https://dict.youdao.com/'},timeout=timeout)
    response.raise_for_status()
    content=response.content
    if len(content)<350 or content[:100].lower().find(b'<html')>=0 or content[:100].lower().find(b'<!doctype')>=0:
        raise ValueError('服务器未返回有效音频（可能是访问限制或单词无发音）')
    Path(dest).write_bytes(content)
    return dest


def ffmpeg_executable():
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except ImportError:
        import shutil
        exe=shutil.which('ffmpeg')
        if exe: return exe
        raise RuntimeError('缺少 FFmpeg。请重新双击「安装依赖.bat」。')


def render_audio(words, output, accent='us', gap_ms=1200, pace=1.0, progress=None, is_cancelled=None, retries=2, repeat=1):
    """Download clips and concatenate in Python/FFmpeg; raise if every clip fails.
    Returns (successful words, failures). Deliberately excludes failures from MP3.
    """
    words=deduplicate(words)
    if not words: raise ValueError('词汇列表为空。')
    if not (1 <= int(repeat) <= 5): raise ValueError('重复次数必须在1至5之间。')
    if not (0 <= int(gap_ms) <= 5000): raise ValueError('停顿应在0至5000毫秒之间。')
    if not (.70 <= float(pace) <= 1.3): raise ValueError('语速范围为0.7–1.3倍。')
    exe=ffmpeg_executable(); good=[]; failures=[]
    output=Path(output); output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='vocab_audio_') as td:
        pieces=[]
        for index,word in enumerate(words,1):
            if is_cancelled and is_cancelled(): raise InterruptedError('用户取消了生成操作')
            if progress: progress(index,len(words),word,'下载中')
            source=Path(td)/f'source_{index:04d}.mp3'; target=Path(td)/f'clip_{index:04d}.mp3'
            problem=None
            for attempt in range(retries+1):
                try:
                    fetch_clip(word,source,accent)
                    # audio filter normalizes levels, optional tempo and silence; use FFmpeg to decode/check source.
                    f=f'aresample=44100,atempo={pace:.3f},loudnorm=I=-20:TP=-2:LRA=10,apad=pad_dur={int(gap_ms)/1000:.3f}'
                    cmd=[exe,'-hide_banner','-loglevel','error','-y','-i',str(source),'-vn','-af',f,'-ar','44100','-ac','1','-codec:a','libmp3lame','-b:a','128k',str(target)]
                    p=subprocess.run(cmd,capture_output=True,text=True,timeout=45,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                    if p.returncode!=0 or target.stat().st_size<500:
                        raise RuntimeError((p.stderr or '转码失败')[-220:])
                    problem=None; break
                except Exception as e:
                    problem=str(e)
                    if attempt<retries: time.sleep(.5*(attempt+1))
            if problem:
                failures.append((word,problem));
                if progress: progress(index,len(words),word,'跳过（下载失败）')
                continue
            pieces.extend([target]*int(repeat)); good.append(word)
            if progress: progress(index,len(words),word,'完成')
            time.sleep(.10) # courteous rate-limiting
        if not pieces: raise RuntimeError('所有单词的音频都下载失败。请检查网络或稍后重试。')
        if is_cancelled and is_cancelled(): raise InterruptedError('用户取消了生成操作')
        listing=Path(td)/'concat.txt'
        listing.write_text(''.join("file '"+str(p).replace("'", "'\\''")+"'\n" for p in pieces),encoding='utf-8')
        final_tmp=Path(td)/'combined.mp3'
        cmd=[exe,'-hide_banner','-loglevel','error','-y','-f','concat','-safe','0','-i',str(listing),'-c:a','libmp3lame','-b:a','128k',str(final_tmp)]
        proc=subprocess.run(cmd,capture_output=True,text=True,timeout=180,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        if proc.returncode!=0: raise RuntimeError('合并失败：'+(proc.stderr or '')[-400:])
        output.write_bytes(final_tmp.read_bytes())
    return good,failures
